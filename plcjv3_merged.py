import json
import random
import time
from collections import Counter
from pathlib import Path

import requests

# --- 用户配置区域 ---
OUTPUT_FILE_TTG = "realtime_matches_history.json"
OUTPUT_FILE_HAD_HHAD = "realtime_matches_had_hhad_history.json"
API_URL = "https://webapi.sporttery.cn/gateway/uniform/football/getFixedBonusV1.qry?clientCode=3001&matchId={}"


def _to_match_id(raw_value):
    try:
        if raw_value is None:
            return None
        return int(str(raw_value).strip())
    except (TypeError, ValueError):
        return None


def _extract_match_ids_from_json(raw_obj):
    match_ids = []
    if isinstance(raw_obj, dict):
        for key, value in raw_obj.items():
            if key in {"matchId", "match_id", "比赛ID", "比赛Id"}:
                match_id = _to_match_id(value)
                if match_id is not None:
                    match_ids.append(match_id)
            if isinstance(value, (dict, list)):
                match_ids.extend(_extract_match_ids_from_json(value))
    elif isinstance(raw_obj, list):
        for item in raw_obj:
            match_ids.extend(_extract_match_ids_from_json(item))
    return match_ids


def _load_match_ids_from_json(json_path):
    try:
        with open(json_path, "r", encoding="utf-8") as f:
            raw_data = json.load(f)
    except Exception:
        return []
    return _extract_match_ids_from_json(raw_data)


def _deduplicate_match_ids(match_ids):
    unique_ids = []
    seen = set()
    for match_id in match_ids:
        if match_id in seen:
            continue
        seen.add(match_id)
        unique_ids.append(match_id)
    return unique_ids


def _load_match_ids_from_xlsx(xlsx_path):
    try:
        from openpyxl import load_workbook
    except Exception:
        return []

    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
    except Exception:
        return []

    id_list = []
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return id_list

    normalized_header = [str(col).strip() if col is not None else "" for col in header]
    target_index = None
    for candidate in ("比赛ID", "matchId", "match_id", "id"):
        if candidate in normalized_header:
            target_index = normalized_header.index(candidate)
            break

    if target_index is None:
        target_index = 0

    for row in rows:
        if not row or target_index >= len(row):
            continue
        match_id = _to_match_id(row[target_index])
        if match_id is not None:
            id_list.append(match_id)

    return id_list


def resolve_match_ids(script_dir=None):
    script_dir = Path(script_dir) if script_dir is not None else Path(__file__).resolve().parent
    candidates = [
        ("json", script_dir / "match_data1.json"),
        ("xlsx", script_dir / "match_data1.xlsx"),
    ]

    for file_type, file_path in candidates:
        if not file_path.exists():
            continue

        if file_type == "json":
            match_ids = _load_match_ids_from_json(file_path)
        else:
            match_ids = _load_match_ids_from_xlsx(file_path)

        match_ids = _deduplicate_match_ids(match_ids)
        if match_ids:
            return match_ids, file_path.name

    raise RuntimeError(
        "未在脚本同目录找到可用比赛ID来源。请提供 match_data1.json 或 match_data1.xlsx，"
        "且其中必须包含 matchId/比赛ID 列。"
    )


def fetch_odds_history(match_id):
    url = API_URL.format(match_id)
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "Referer": "https://www.sporttery.cn/",
    }

    try:
        print(f"[{match_id}] 正在连接...", end="")
        response = requests.get(url, headers=headers, timeout=8)
        if response.status_code != 200:
            print(f" ❌ HTTP {response.status_code}")
            return None

        data = response.json()
        value = data.get("value")
        if not value:
            print(" ❌ value为空 (无效ID)")
            return None

        odds_history = value.get("oddsHistory")
        if not odds_history:
            print(" ⚠️ 无oddsHistory (未开售)")
            return None
        return odds_history
    except Exception as e:
        print(f" ❌ 异常: {e}")
        return None


def build_ttg_match_data(match_id, odds_history):
    ttg_list = odds_history.get("ttgList", [])
    if not ttg_list:
        return None

    def extract_odds(raw_data):
        try:
            return {
                "g0": float(raw_data.get("s0", 0)),
                "g1": float(raw_data.get("s1", 0)),
                "g2": float(raw_data.get("s2", 0)),
                "g3": float(raw_data.get("s3", 0)),
                "g4": float(raw_data.get("s4", 0)),
                "g5": float(raw_data.get("s5", 0)),
                "g6": float(raw_data.get("s6", 0)),
                "g7plus": float(raw_data.get("s7", 0)),
                "time": raw_data.get("updateTime", ""),
            }
        except Exception:
            return None

    timeline = []
    min_odds = [999.0] * 8
    max_odds = [-1.0] * 8
    keys_map = ["g0", "g1", "g2", "g3", "g4", "g5", "g6", "g7plus"]

    for raw_item in ttg_list:
        clean_item = extract_odds(raw_item)
        if not clean_item:
            continue
        timeline.append(clean_item)
        for idx, key in enumerate(keys_map):
            val = clean_item.get(key, 0)
            if val > 0:
                if val < min_odds[idx]:
                    min_odds[idx] = val
                if val > max_odds[idx]:
                    max_odds[idx] = val

    if not timeline:
        return None

    odds_initial = timeline[0]
    odds_final = timeline[-1]
    change_count = len(timeline)

    trend_symbol = "➡️"
    diff = odds_final["g0"] - odds_initial["g0"]
    if diff > 0.5:
        trend_symbol = "⬆️"
    elif diff < -0.5:
        trend_symbol = "⬇️"

    volatility_alert = ""
    if max_odds[0] > max(odds_initial["g0"], odds_final["g0"]) + 1.0:
        volatility_alert = "⚡(曾冲高)"
    elif min_odds[0] < min(odds_initial["g0"], odds_final["g0"]) - 1.0:
        volatility_alert = "⚡(曾砸盘)"

    league_name = odds_history.get("leagueAbbName", "未知联赛")
    home = odds_history.get("homeTeamAbbName", "主队")
    away = odds_history.get("awayTeamAbbName", "客队")

    return {
        "match_id": match_id,
        "display_name": f"{league_name}: {home} vs {away}",
        "league": league_name,
        "odds": {
            "g0": odds_final["g0"],
            "g1": odds_final["g1"],
            "g2": odds_final["g2"],
            "g3": odds_final["g3"],
            "g4": odds_final["g4"],
            "g5": odds_final["g5"],
            "g6": odds_final["g6"],
            "g7plus": odds_final["g7plus"],
            "history": {
                "initial": odds_initial,
                "final": odds_final,
                "timeline": timeline,
                "stats": {
                    "count": change_count,
                    "trend": trend_symbol,
                    "volatility": volatility_alert,
                    "min_odds": min_odds,
                    "max_odds": max_odds,
                },
            },
        },
    }


def _calc_history_stats(timeline, keys_map, primary_key):
    min_odds = [999.0] * len(keys_map)
    max_odds = [-1.0] * len(keys_map)

    for item in timeline:
        for idx, key in enumerate(keys_map):
            val = item.get(key, 0)
            if isinstance(val, (int, float)) and val > 0:
                if val < min_odds[idx]:
                    min_odds[idx] = val
                if val > max_odds[idx]:
                    max_odds[idx] = val

    odds_initial = timeline[0]
    odds_final = timeline[-1]
    change_count = 0
    for prev, curr in zip(timeline, timeline[1:]):
        changed = False
        for k in keys_map:
            pv = prev.get(k, 0)
            cv = curr.get(k, 0)
            if isinstance(pv, (int, float)) and isinstance(cv, (int, float)):
                if abs(pv - cv) > 1e-12:
                    changed = True
                    break
        if changed:
            change_count += 1

    def _trend_for(key):
        diff = odds_final.get(key, 0) - odds_initial.get(key, 0)
        if diff > 0.5:
            return "⬆️"
        if diff < -0.5:
            return "⬇️"
        return "➡️"

    def _volatility_for(key, idx):
        initial_v = odds_initial.get(key, 0)
        final_v = odds_final.get(key, 0)
        if max_odds[idx] > max(initial_v, final_v) + 1.0:
            return "⚡(曾冲高)"
        if min_odds[idx] < min(initial_v, final_v) - 1.0:
            return "⚡(曾砸盘)"
        return ""

    trend_map = {k: _trend_for(k) for k in keys_map}
    volatility_map = {k: _volatility_for(k, i) for i, k in enumerate(keys_map)}
    trend_symbol = trend_map.get(primary_key, "➡️")
    volatility_alert = volatility_map.get(primary_key, "")

    return (
        odds_initial,
        odds_final,
        change_count,
        trend_symbol,
        volatility_alert,
        min_odds,
        max_odds,
        trend_map,
        volatility_map,
    )


def build_had_hhad_match_data(match_id, odds_history):
    had_list = odds_history.get("hadList", [])
    hhad_list = odds_history.get("hhadList", [])

    if not had_list and not hhad_list:
        return None

    def extract_1x2(raw_data, keep_goal_line=False):
        try:
            update_date = raw_data.get("updateDate", "") or ""
            update_time = raw_data.get("updateTime", "") or ""
            dt = f"{update_date} {update_time}".strip()
            item = {
                "h": float(raw_data.get("h", 0)),
                "d": float(raw_data.get("d", 0)),
                "a": float(raw_data.get("a", 0)),
                "time": dt,
            }
            if keep_goal_line:
                item["goalLine"] = raw_data.get("goalLine", "")
            return item
        except (TypeError, ValueError):
            return None

    keys_map = ["h", "d", "a"]
    had_timeline = [extract_1x2(x, keep_goal_line=False) for x in had_list]
    had_timeline = [x for x in had_timeline if x]
    had_timeline.sort(key=lambda x: x.get("time") or "9999-99-99 99:99:99")

    hhad_timeline_raw = [extract_1x2(x, keep_goal_line=True) for x in hhad_list]
    hhad_timeline_raw = [x for x in hhad_timeline_raw if x]

    hhad_timelines_by_gl = {}
    for item in hhad_timeline_raw:
        gl = item.get("goalLine", "") or ""
        hhad_timelines_by_gl.setdefault(gl, []).append(item)

    for tl in hhad_timelines_by_gl.values():
        tl.sort(key=lambda x: x.get("time") or "9999-99-99 99:99:99")

    main_goal_line = ""
    if hhad_timeline_raw:
        goal_lines = [x.get("goalLine", "") or "" for x in hhad_timeline_raw]
        main_goal_line = Counter(goal_lines).most_common(1)[0][0]
    hhad_timeline = hhad_timelines_by_gl.get(main_goal_line, [])

    if had_list and not had_timeline:
        return None
    if hhad_list and not hhad_timeline_raw:
        return None

    had_section = None
    hhad_section = None
    hhad_by_goalLine_section = None

    if had_timeline:
        (
            had_initial,
            had_final,
            had_change_count,
            had_trend_symbol,
            had_volatility,
            had_min_odds,
            had_max_odds,
            had_trend_all,
            had_volatility_all,
        ) = _calc_history_stats(had_timeline, keys_map, primary_key="h")
        had_section = {
            "h": had_final["h"],
            "d": had_final["d"],
            "a": had_final["a"],
            "history": {
                "initial": had_initial,
                "final": had_final,
                "timeline": had_timeline,
                "stats": {
                    "count": had_change_count,
                    "trend": had_trend_symbol,
                    "volatility": had_volatility,
                    "trend_all": had_trend_all,
                    "volatility_all": had_volatility_all,
                    "min_odds": had_min_odds,
                    "max_odds": had_max_odds,
                },
            },
        }

    if hhad_timeline_raw:
        if hhad_timeline:
            (
                hhad_initial,
                hhad_final,
                hhad_change_count,
                hhad_trend_symbol,
                hhad_volatility,
                hhad_min_odds,
                hhad_max_odds,
                hhad_trend_all,
                hhad_volatility_all,
            ) = _calc_history_stats(hhad_timeline, keys_map, primary_key="h")
            hhad_section = {
                "goalLine": main_goal_line,
                "h": hhad_final["h"],
                "d": hhad_final["d"],
                "a": hhad_final["a"],
                "history": {
                    "initial": hhad_initial,
                    "final": hhad_final,
                    "timeline": hhad_timeline,
                    "stats": {
                        "count": hhad_change_count,
                        "trend": hhad_trend_symbol,
                        "volatility": hhad_volatility,
                        "trend_all": hhad_trend_all,
                        "volatility_all": hhad_volatility_all,
                        "min_odds": hhad_min_odds,
                        "max_odds": hhad_max_odds,
                    },
                },
            }

        hhad_by_goalLine_section = {}
        for gl, tl in hhad_timelines_by_gl.items():
            if not tl:
                continue
            (
                gi,
                gf,
                gc,
                gtrend,
                gvol,
                gmin,
                gmax,
                gall_trend,
                gall_vol,
            ) = _calc_history_stats(tl, keys_map, primary_key="h")
            hhad_by_goalLine_section[gl] = {
                "goalLine": gl,
                "h": gf["h"],
                "d": gf["d"],
                "a": gf["a"],
                "history": {
                    "initial": gi,
                    "final": gf,
                    "timeline": tl,
                    "stats": {
                        "count": gc,
                        "trend": gtrend,
                        "volatility": gvol,
                        "trend_all": gall_trend,
                        "volatility_all": gall_vol,
                        "min_odds": gmin,
                        "max_odds": gmax,
                    },
                },
            }

    league_name = odds_history.get("leagueAbbName", "未知联赛")
    home = odds_history.get("homeTeamAbbName", "主队")
    away = odds_history.get("awayTeamAbbName", "客队")

    odds_obj = {}
    if had_section is not None:
        odds_obj["had"] = had_section
    if hhad_section is not None:
        odds_obj["hhad"] = hhad_section
    if hhad_by_goalLine_section is not None:
        odds_obj["hhad_by_goalLine"] = hhad_by_goalLine_section

    if not odds_obj:
        return None

    return {
        "match_id": match_id,
        "display_name": f"{league_name}: {home} vs {away}",
        "league": league_name,
        "odds": odds_obj,
    }


def main(match_ids=None, id_source_file=None):
    if match_ids is None:
        match_ids, id_source_file = resolve_match_ids()
    else:
        match_ids = _deduplicate_match_ids(match_ids)
        id_source_file = id_source_file or "传入ID列表"

    if not match_ids:
        raise RuntimeError("没有可采集的比赛ID。")

    ttg_matches = []
    had_hhad_matches = []

    id_min = min(match_ids)
    id_max = max(match_ids)

    print(
        f"=== 启动合并采集器 V1.1 (ID范围: {id_min} - {id_max} | 来源: {id_source_file} | 实际采集ID数: {len(match_ids)}) ===\n"
    )

    for mid in match_ids:
        odds_history = fetch_odds_history(mid)
        if odds_history:
            ttg_data = build_ttg_match_data(mid, odds_history)
            had_hhad_data = build_had_hhad_match_data(mid, odds_history)

            league_name = odds_history.get("leagueAbbName", "未知联赛")
            home = odds_history.get("homeTeamAbbName", "主队")
            away = odds_history.get("awayTeamAbbName", "客队")
            tags = []
            if ttg_data:
                tags.append("TTG")
                ttg_matches.append(ttg_data)
            if had_hhad_data:
                tags.append("HAD/HHAD")
                had_hhad_matches.append(had_hhad_data)
            print(f" ✅ {league_name} | {home} vs {away} ({' + '.join(tags) if tags else '无可用玩法'})")

        time.sleep(random.uniform(0.5, 1.0))

    print("-" * 50)
    print(f"采集结束：TTG有效 {len(ttg_matches)} 场，HAD/HHAD有效 {len(had_hhad_matches)} 场。")

    ttg_match_ids = {item["match_id"] for item in ttg_matches}
    had_hhad_match_ids = {item["match_id"] for item in had_hhad_matches}
    missing_ttg_ids = [mid for mid in match_ids if mid not in ttg_match_ids]
    missing_had_hhad_ids = [mid for mid in match_ids if mid not in had_hhad_match_ids]
    if missing_ttg_ids:
        print(f"TTG未产出ID: {', '.join(map(str, missing_ttg_ids))}")
    if missing_had_hhad_ids:
        print(f"HAD/HHAD未产出ID: {', '.join(map(str, missing_had_hhad_ids))}")

    if ttg_matches:
        with open(OUTPUT_FILE_TTG, "w", encoding="utf-8") as f:
            json.dump(ttg_matches, f, ensure_ascii=False, indent=4)
        print(f"TTG数据文件已生成: {OUTPUT_FILE_TTG}")
    else:
        print("TTG未获取到数据。")

    if had_hhad_matches:
        with open(OUTPUT_FILE_HAD_HHAD, "w", encoding="utf-8") as f:
            json.dump(had_hhad_matches, f, ensure_ascii=False, indent=4)
        print(f"HAD/HHAD数据文件已生成: {OUTPUT_FILE_HAD_HHAD}")
    else:
        print("HAD/HHAD未获取到数据。")


if __name__ == "__main__":
    main()
